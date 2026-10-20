import json
import re
import glob
import os
import sys

CARDS_DIR = "D:/projects/fgo/Scripts/Cards"
POWERS_DIR = "D:/projects/fgo/Scripts/Powers"
ZHS_PATH = "D:/projects/fgo/Fgo/localization/zhs/cards.json"
ENG_PATH = "D:/projects/fgo/Fgo/localization/eng/cards.json"
JPN_PATH = "D:/projects/fgo/Fgo/localization/jpn/cards.json"
OUT_PATH = "D:/projects/fgo/cards_export.xlsx"

SKIP_CLASSES = {"FgoCardModel", "FgoCooldownCardModel", "FgoBaseCardModel", "ModCardTemplate",
                "NobleCardModel", "FgoDishCardModel"}

RARITY_ZH = {
    "Basic": "初始",
    "Common": "普通",
    "Uncommon": "罕见",
    "Rare": "稀有",
    "Ancient": "先古",
}

RARITY_RANK = {"初始": 1, "普通": 2, "罕见": 3, "稀有": 4, "先古": 5}

TYPE_EN = {
    "Attack": "ATTACK",
    "Skill": "SKILL",
    "Power": "POWER",
    "Status": "STATUS",
    "Curse": "CURSE",
}

PLACEHOLDER_FEEDBACK = "看到这个请向作者反馈此卡牌缺少本地化"


# ---------------------------------------------------------------- constants


def build_global_const_table():
    """Collect `public const T Name = V` from Cards+Powers sources as ClassName.Name -> V."""
    table = {}
    paths = glob.glob(os.path.join(CARDS_DIR, "**", "*.cs"), recursive=True)
    paths += glob.glob(os.path.join(POWERS_DIR, "*.cs"))
    for path in paths:
        src = open(path, encoding="utf-8").read()
        mcls = re.search(r"public\s+(?:abstract\s+|sealed\s+|static\s+)?class\s+(\w+)", src)
        if not mcls:
            continue
        cls = mcls.group(1)
        for m in re.finditer(
            r"public\s+const\s+(?:int|long|decimal|float|double|short)\s+(\w+)\s*=\s*(-?[\d.]+)m?",
            src,
        ):
            name, val = m.group(1), m.group(2)
            v = float(val)
            table[f"{cls}.{name}"] = int(v) if v.is_integer() else v
    return table


def parse_local_consts(src):
    """All const declarations + override expression-bodied props in this file."""
    consts = {}
    for m in re.finditer(
        r"const\s+(?:int|long|decimal|float|double|short)\s+(\w+)\s*=\s*(-?[\d.]+)m?",
        src,
    ):
        name, val = m.group(1), m.group(2)
        v = float(val)
        consts[name] = int(v) if v.is_integer() else v
    for m in re.finditer(
        r"const\s+string\s+(\w+)\s*=\s*\"([^\"]+)\"",
        src,
    ):
        consts[m.group(1)] = m.group(2)
    for m in re.finditer(
        r"public\s+override\s+(?:int|long|decimal)\s+(\w+)\s*=>\s*(-?[\d.]+)m?\s*;",
        src,
    ):
        name, val = m.group(1), m.group(2)
        v = float(val)
        consts[name] = int(v) if v.is_integer() else v
    return consts


def eval_expr(expr, consts):
    """Evaluate a C#-ish numeric expression: literals, const names, Class.Const, + - * /."""
    if expr is None:
        return None
    e = expr.strip()
    if not e:
        return None
    if re.fullmatch(r"-?\d+(?:\.\d+)?m?", e):
        v = float(e.rstrip("m"))
        return int(v) if v.is_integer() else v
    idents = set(re.findall(r"[A-Za-z_]\w*(?:\.[A-Za-z_]\w*)*", e))
    resolved = e
    for ident in sorted(idents, key=len, reverse=True):
        if ident in consts:
            resolved = re.sub(r"\b" + re.escape(ident) + r"\b", repr(consts[ident]), resolved)
        else:
            return None
    if not re.fullmatch(r"[-+*/(). \d.e]+", resolved):
        return None
    try:
        v = eval(resolved, {"__builtins__": {}}, {})
    except Exception:
        return None
    if not isinstance(v, (int, float)):
        return None
    return int(v) if float(v).is_integer() else v


def split_args(s):
    """Split a C# argument list on top-level commas (handles nesting, strings, =>)."""
    parts = []
    depth = 0
    cur = []
    in_str = False
    i = 0
    while i < len(s):
        c = s[i]
        if in_str:
            cur.append(c)
            if c == "\\" and i + 1 < len(s):
                cur.append(s[i + 1])
                i += 2
                continue
            if c == '"':
                in_str = False
            i += 1
            continue
        if c == '"':
            in_str = True
            cur.append(c)
            i += 1
            continue
        if c == "=" and i + 1 < len(s) and s[i + 1] == ">":
            cur.append("=>")
            i += 2
            continue
        if c in "([{":
            depth += 1
        elif c in ")]}":
            depth -= 1
        if c == "," and depth == 0:
            parts.append("".join(cur))
            cur = []
        else:
            cur.append(c)
        i += 1
    if cur:
        parts.append("".join(cur))
    return [p.strip() for p in parts]


# ---------------------------------------------------------------- parsing


def extract_balanced(text, start):
    """From text[start] which must be '(' or '[', return the balanced substring including the bracket."""
    open_ch = text[start]
    close_ch = ")" if open_ch == "(" else "]"
    depth = 0
    i = start
    n = len(text)
    while i < n:
        c = text[i]
        if c == open_ch:
            depth += 1
        elif c == close_ch:
            depth -= 1
            if depth == 0:
                return text[start : i + 1]
        elif c == '"' or c == "'":
            quote = c
            i += 1
            while i < n and text[i] != quote:
                if text[i] == "\\":
                    i += 1
                i += 1
        i += 1
    return text[start:]


def parse_card_file(path, global_consts):
    src = open(path, encoding="utf-8").read()
    m = re.search(r"public\s+(?:abstract\s+|sealed\s+)?class\s+(\w+)", src)
    if not m:
        return None
    cls = m.group(1)
    mb = re.search(r":\s*\w+\s*\(", src)
    cost = None
    rarity = None
    card_type = None
    if mb:
        args_block = extract_balanced(src, mb.end() - 1)
        mi = re.search(r"-?\d+", args_block)
        if mi:
            cost = int(mi.group(0))
        mr = re.search(r"CardRarity\.(\w+)", args_block)
        if mr:
            rarity = mr.group(1)
        mt = re.search(r"CardType\.(\w+)", args_block)
        if mt:
            card_type = mt.group(1)
    consts = parse_local_consts(src)
    consts.update({k: v for k, v in global_consts.items()})
    vars_dict = parse_vars(src, consts)
    upgrades = parse_upgrades(src, consts)
    # fallback: override props (e.g. CooldownMax => 12) fill card-face vars not declared
    # in CanonicalVars but referenced by the description
    for k, v in consts.items():
        if k not in vars_dict and "." not in k:
            vars_dict.setdefault(k, v)
    # fallback: PowerCmd.Apply<XxxPower>(..., <const expr>, ...) amounts for cards whose
    # CanonicalVars are missing (e.g. EightKindness applies `const int amount = 1`)
    for pm in re.finditer(r"PowerCmd\.Apply<(\w+)>\s*\(", src):
        pname = pm.group(1)
        args_block = extract_balanced(src, pm.end() - 1)
        args = split_args(args_block[1:-1])
        if len(args) >= 3:
            val = eval_expr(args[2], consts)
            if val is not None and pname not in vars_dict:
                vars_dict[pname] = val
    return {
        "class": cls,
        "cost": cost,
        "rarity": rarity,
        "type": card_type,
        "vars": vars_dict,
        "upgrades": upgrades,
    }


def parse_vars(src, consts):
    """Extract all ModCardVars.Xxx(...) calls from CanonicalVars / AdditionalCanonicalVars blocks."""
    vars_dict = {}
    for kw in ("CanonicalVars", "AdditionalCanonicalVars"):
        idx = src.find(kw)
        while idx != -1:
            j = src.find("[", idx)
            k = src.find("(", idx)
            if j == -1 and k == -1:
                break
            if j != -1 and (k == -1 or j < k):
                block = extract_balanced(src, j)
            else:
                block = extract_balanced(src, k)
            for name, value in find_modcardvars_calls(block, consts):
                if name and value is not None:
                    vars_dict[name] = value
            idx = src.find(kw, idx + len(kw))
    return vars_dict


def parse_upgrades(src, consts):
    """Extract upgrade deltas from OnUpgrade: DynamicVars.X / ["X"] / [nameof(X)].UpgradeValueBy(expr).

    The argument may be a literal (optionally with a `m` suffix) or a const expression
    (e.g. `UpgradeStacks * BlessingOfKurPower.MaxHpPerStack`). A call is "power-prefixed"
    when its `DynamicVars` is accessed through a member/power object (e.g.
    `power.DynamicVars[...]` or `_fifthForcePower?.DynamicVars[...]`); that upgrades the
    power's hover stat, not the card description. Resolution prefers the card-level delta;
    falls back to the power-prefixed one only when no card-level upgrade of that name exists.
    """
    card_deltas = {}
    power_deltas = {}
    arg = r"((?:[^()]|\([^()]*\))*)"
    pats = [
        r"DynamicVars\s*\.\s*(\w+)\s*\.\s*UpgradeValueBy\(\s*" + arg + r"\s*\)",
        r'DynamicVars\s*\[\s*"([^"]+)"\s*\]\s*\.\s*UpgradeValueBy\(\s*' + arg + r"\s*\)",
        r"DynamicVars\s*\[\s*nameof\(\s*(\w+)\s*\)\s*\]\s*\.\s*UpgradeValueBy\(\s*" + arg + r"\s*\)",
    ]
    for pat in pats:
        for m in re.finditer(pat, src):
            name = m.group(1)
            n = eval_expr(m.group(2), consts)
            if n is None:
                continue
            ctx = src[max(0, m.start() - 12) : m.start()]
            power = bool(re.search(r"[\w)\]](?:\?)?\.\s*$", ctx))
            target = power_deltas if power else card_deltas
            target[name] = target.get(name, 0) + n
    deltas = {}
    deltas.update(power_deltas)  # fallback: power-prefixed only
    deltas.update(card_deltas)   # card-level takes priority
    return deltas


def find_modcardvars_calls(block, consts):
    calls = []
    i = block.find("ModCardVars.")
    while i != -1:
        p = block.find("(", i)
        if p == -1:
            break
        args = extract_balanced(block, p)
        method = block[i + len("ModCardVars.") : p]
        name, value = interpret_call(method, args, consts)
        if name:
            calls.append((name, value))
        i = block.find("ModCardVars.", p)
    return calls


def interpret_call(method, args, consts):
    # args includes parentheses
    inner = args[1:-1]
    base_method = re.sub(r"<[^>]*>", "", method)
    parts = split_args(inner)
    # Power<XxxPower>
    if base_method == "Power":
        mg = re.search(r"<(\w+)>", method)
        name = mg.group(1) if mg else None
        value = eval_expr(parts[0], consts) if parts else None
        return name, value
    if base_method in ("Int", "Computed") or base_method.startswith("Computed"):
        name = None
        value = None
        if parts:
            first = parts[0].strip()
            # only a bare string literal counts as the var name; a lambda body may
            # coincidentally contain string literals (ComputedPowerAmountGiven etc.)
            sm = re.fullmatch(r'"([^"]+)"', first)
            if sm:
                name = sm.group(1)
            elif re.fullmatch(r"\w+", first) and first in consts:
                name = str(consts[first])  # const string key, e.g. FatigueKey
        # value: first remaining arg that evaluates (skips lambdas / DynamicVars refs)
        for arg in parts[1:]:
            v = eval_expr(arg, consts)
            if v is not None:
                value = v
                break
        return name, value
    # typed vars: Damage, Block, Heal, Cards, Energy, HpLoss, Strength, ...
    name = base_method
    value = eval_expr(parts[0], consts) if parts else None
    return name, value


# ---------------------------------------------------------------- rendering


def handle_choose(text, vars_dict):
    """{Var:choose(N):b0|b1|...} -> pick branch by value (clamped), N as fallback index."""
    out = []
    i = 0
    while True:
        m = re.search(r"\{(\w+):choose\((-?\d+)\):", text[i:])
        if not m:
            out.append(text[i:])
            break
        out.append(text[i : i + m.start()])
        var = m.group(1)
        fb = int(m.group(2))
        j = i + m.end()
        depth = 1
        branches = []
        cur = []
        k = j
        while k < len(text) and depth > 0:
            c = text[k]
            if c == "{":
                depth += 1
            elif c == "}":
                depth -= 1
                if depth == 0:
                    break
            if c == "|" and depth == 1:
                branches.append("".join(cur))
                cur = []
            else:
                cur.append(c)
            k += 1
        branches.append("".join(cur))
        val = vars_dict.get(var)
        if val is None:
            val = fb
        try:
            val = int(val)
        except (TypeError, ValueError):
            val = fb
        idx = max(0, min(val, len(branches) - 1))
        out.append(branches[idx])
        i = k + 1
    return "".join(out)


def resolve(text, vars_dict, cost, upgrades=None):
    if upgrades is None:
        upgrades = {}
    if text is None:
        return None
    # strip style tags
    text = re.sub(r"\[gold\](.*?)\[/gold\]", r"\1", text, flags=re.S)
    text = re.sub(r"\[purple\](.*?)\[/purple\]", r"\1", text, flags=re.S)

    def fmt(name):
        base = vars_dict.get(name)
        if base is None:
            return "{" + name + "}"
        delta = upgrades.get(name, 0)
        if delta:
            return f"{base}[{base + delta}]"
        return str(base)

    # energy icons token: {energyPrefix:energyIcons(3)} or {Energy:energyIcons()}
    def repl_energy(m):
        var = m.group(1)
        num = m.group(2)
        if num:
            return f"{int(num)}能量"
        return fmt(var) + "能量"

    text = re.sub(r"\{(\w+):[^{}]*energyIcons\((\d*)\)[^{}]*\}", repl_energy, text)

    # conditionals: repeat until stable (branches may nest)
    prev = None
    while prev != text:
        prev = text
        # IfUpgraded:show:X|Y -> base Y, plus upgraded-only X in 〔〕
        def repl_ifup(m):
            x = m.group(1)
            y = m.group(2)
            return y + ("〔" + x + "〕" if x.strip() else "")

        text = re.sub(r"\{IfUpgraded:show:(.*?)\|(.*?)\}", repl_ifup, text, flags=re.S)
        # InCombat:X|Y -> default (non-combat) = Y
        text = re.sub(
            r"\{InCombat:(.*?)\|(.*?)\}",
            lambda m: m.group(2),
            text,
            flags=re.S,
        )
        # choose: {Var:choose(N):b0|b1|...}
        text = handle_choose(text, vars_dict)
        # plural: {Var:plural:singular|plural} -> noun by count (base==1 -> singular)
        def repl_plural(m):
            var = m.group(1)
            sing = m.group(2)
            plu = m.group(3)
            return sing if vars_dict.get(var) == 1 else plu

        text = re.sub(
            r"\{(\w+):plural:([^}|]*)\|([^}]*)\}",
            repl_plural,
            text,
            flags=re.S,
        )
        # cond: {Var:cond:=1?TRUE|FALSE}  (FALSE may be {:diff()} with a nested brace)
        def repl_cond_diff(m):
            var = m.group(1)
            op = m.group(2)
            tru = m.group(3)
            val = vars_dict.get(var)
            if op.startswith("="):
                try:
                    if val == int(op[1:]):
                        return tru
                except ValueError:
                    pass
            return fmt(var)

        def repl_cond_lit(m):
            var = m.group(1)
            op = m.group(2)
            tru = m.group(3)
            fal = m.group(4)
            val = vars_dict.get(var)
            if op.startswith("="):
                try:
                    if val == int(op[1:]):
                        return tru
                except ValueError:
                    pass
            return fal

        text = re.sub(
            r"\{(\w+):cond:(=[0-9]+)\?([^{}|]*)\|\{:diff\(\)\}\}",
            repl_cond_diff,
            text,
            flags=re.S,
        )
        text = re.sub(
            r"\{(\w+):cond:(=[0-9]+)\?([^{}|]*)\|([^{}]*)\}",
            repl_cond_lit,
            text,
            flags=re.S,
        )

    # diff / plain var resolution
    def repl_var(m):
        var = m.group(1)
        return fmt(var)

    prev = None
    while prev != text:
        prev = text
        text = re.sub(r"\{(\w+):diff\(\)\}", repl_var, text)
        text = re.sub(r"\{(\w+):inverseDiff\(\)\}", repl_var, text)
        text = re.sub(r"\{(\w+)\}", repl_var, text)

    # cleanup leftover empty braces and spaces
    text = text.replace("{}", "")
    return text


def card_key(cls):
    return "FGO_CARD_" + re.sub(r"(?<!^)(?=[A-Z])", "_", cls).upper()


def collect_rows(global_consts):
    zhs = json.load(open(ZHS_PATH, encoding="utf-8"))
    eng = json.load(open(ENG_PATH, encoding="utf-8"))
    jpn = json.load(open(JPN_PATH, encoding="utf-8"))

    rows = []
    unresolved_report = []
    for path in sorted(glob.glob(os.path.join(CARDS_DIR, "*.cs"))):
        info = parse_card_file(path, global_consts)
        if not info:
            continue
        if info["class"] in SKIP_CLASSES:
            continue
        key = card_key(info["class"])
        ztitle = zhs.get(key + ".title")
        zdesc = zhs.get(key + ".description")
        etitle = eng.get(key + ".title")
        edesc = eng.get(key + ".description")
        jtitle = jpn.get(key + ".title")
        jdesc = jpn.get(key + ".description")
        # skip placeholder feedback entries (abstract models / missing localization)
        if ztitle == PLACEHOLDER_FEEDBACK or (ztitle is None and zdesc is None):
            unresolved_report.append((info["class"], "no zhs localization (key=" + key + ")"))
            continue
        if jtitle is None or jdesc is None:
            unresolved_report.append((info["class"], f"missing jpn title/desc (key={key})"))
        if etitle is None or edesc is None:
            unresolved_report.append((info["class"], f"missing eng title/desc (key={key})"))
        upg = info.get("upgrades", {})
        r_zh = resolve(zdesc, info["vars"], info["cost"], upg)
        r_en = resolve(edesc, info["vars"], info["cost"], upg)
        r_jp = resolve(jdesc, info["vars"], info["cost"], upg)
        for lang, resolved in (("zh", r_zh), ("en", r_en), ("jp", r_jp)):
            if resolved:
                toks = re.findall(r"\{(\w+)\}", resolved)
                if toks:
                    unresolved_report.append(
                        (info["class"], f"{lang} unresolved: {sorted(set(toks))}")
                    )
        rarity_zh = RARITY_ZH.get(info["rarity"], info["rarity"] or "")
        type_en = TYPE_EN.get(info["type"], info["type"] or "")
        rows.append(
            {
                "key": key,
                "name_zh": ztitle,
                "desc_zh": r_zh,
                "cost": info["cost"],
                "rarity": rarity_zh,
                "type": type_en,
                "name_en": etitle,
                "desc_en": r_en,
                "name_jp": jtitle,
                "desc_jp": r_jp,
            }
        )

    rows.sort(
        key=lambda r: (
            RARITY_RANK.get(r["rarity"], 9),
            r["cost"] if r["cost"] is not None else 0,
            r["name_zh"] or "",
        )
    )
    return rows, unresolved_report


def main():
    global_consts = build_global_const_table()
    rows, unresolved_report = collect_rows(global_consts)

    # local xlsx is optional: the Tencent Doc is the source of truth
    if "--xlsx" in sys.argv:
        write_xlsx(rows)

    print(f"卡牌行数：{len(rows)}")
    print(f"未解析项：{len(unresolved_report)}")
    for item in unresolved_report:
        print("  -", item[0], ":", item[1])
    return rows, unresolved_report


def write_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "卡牌"
    headers = [
        "卡名", "描述", "费用", "稀有度", "类型",
        "英文卡名(title)", "英文描述(description)",
        "日文卡名(title)", "日文描述(description)",
    ]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="4472C4")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", horizontal="center")
    for r in rows:
        ws.append(
            [
                r["name_zh"], r["desc_zh"], r["cost"], r["rarity"], r["type"],
                r["name_en"], r["desc_en"], r["name_jp"], r["desc_jp"],
            ]
        )
    widths = [22, 60, 8, 10, 10, 28, 70, 28, 70]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    wb.save(OUT_PATH)
    print(f"导出完成：{OUT_PATH}")


if __name__ == "__main__":
    main()
