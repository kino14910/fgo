import sys
import json
from openpyxl import load_workbook

TD = "C:/Users/14910/.workbuddy/plugins/cache/workbuddy-builtin/tencent-docs-plugin/5.5.3-wb.37748631.g104760a2.h1a8f7c37fe76/skills/tencent-docs"
sys.path.insert(0, TD)
import tencentdocs as td

FILE_URL = "https://docs.qq.com/sheet/DSGR3bE1zZkJoR1Jh"
XLSX = "D:/projects/fgo/cards_export.xlsx"
SHEET_NAME = "卡牌"

# 1. read local xlsx
wb = load_workbook(XLSX)
ws = wb.active
grid = [list(r) for r in ws.iter_rows(values_only=True)]
print("local rows:", len(grid), "cols:", len(grid[0]))

# 2. 子表「卡牌」已在上一轮创建（空表，ID 见下）；直接复用，避免重名
new_id = "qbPkXC"
print("reuse sheet_id:", new_id)


# 3. build values
values = []
for r, row in enumerate(grid):
    for c, cell in enumerate(row):
        if cell is None:
            cell = ""
        if c == 2 and r >= 1:  # 费用列(数据行) -> NUMBER；表头行保持文本
            try:
                num = int(cell)
            except (ValueError, TypeError):
                num = 0
            values.append({"row": r, "col": c, "value_type": "NUMBER", "number_value": num})
        else:
            values.append({"row": r, "col": c, "value_type": "STRING", "string_value": str(cell)})
print("values cells:", len(values))

# 4. write
res2, err2 = td.call_tool("sheet-mcp", "set_range_value", {
    "file_url": FILE_URL,
    "sheet_id": new_id,
    "values": values,
})
print("set_range_value err:", err2)
print("set_range_value result:", json.dumps(res2, ensure_ascii=False)[:300])

# 5. verify a small range
res3, err3 = td.call_tool("sheet-mcp", "get_cell_data", {
    "file_url": FILE_URL,
    "sheet_id": new_id,
    "start_row": 0, "start_col": 0, "end_row": 5, "end_col": 5,
    "return_csv": True,
})
print("verify err:", err3)
if res3:
    txt = json.dumps(res3, ensure_ascii=False)
    print("verify sample:", txt[:400])

